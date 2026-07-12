"""XLSX-Export: Struktur- und Semantik-Smoke-Test."""

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from kassensturz.model import Transaction
from kassensturz.report.xlsx import _detect_abos, build_workbook
from kassensturz.store import Store


def tx(**kwargs) -> Transaction:
    defaults = dict(
        booking_date=date(2024, 6, 1),
        amount=Decimal("-10.00"),
        currency="EUR",
        account="test",
        counterparty="Shop",
        purpose="Kauf",
    )
    defaults.update(kwargs)
    return Transaction(**defaults)


def test_workbook_sheets_and_consolidation(tmp_path, monkeypatch):
    # neutrales cwd: sonst zieht build_workbook Marcs echtes budget.yaml
    monkeypatch.chdir(tmp_path)
    store = Store(":memory:")
    store.add(
        [
            tx(booking_date=date(2024, 3, 1), amount=Decimal("-100.00"),
               category="Shopping"),
            # Retoure: mindert die Shopping-Ausgaben, ist KEIN Einkommen
            tx(booking_date=date(2024, 3, 9), amount=Decimal("30.00"),
               category="Shopping", purpose="retoure"),
            tx(booking_date=date(2024, 4, 1), amount=Decimal("3000.00"),
               category="Gehalt", purpose="lohn"),
            # neutral & Sondereffekt dürfen NICHT in Einnahmen/Ausgaben zählen
            tx(booking_date=date(2024, 5, 1), amount=Decimal("-500.00"),
               category="Kreditkarten-Abrechnung"),
            tx(booking_date=date(2024, 5, 2), amount=Decimal("9999.00"),
               category="Exit (Sondereffekt)"),
            tx(booking_date=date(2025, 1, 5), amount=Decimal("-40.00"),
               category=None, purpose="tail"),
            # aktives Abo — sonst bleiben Verträge-KPI-Zweige unausgeführt
            # (genau so schlüpfte der mapping-NameError an den Tests vorbei)
            *[tx(booking_date=date(2025, m, 3), amount=Decimal("-9.99"),
                 counterparty="Netflix", category="Abos & Dienste")
              for m in range(1, 7)],
        ]
    )
    store.record_balances([("test", "2025-06-01", 123456, "EUR", "t")])
    out = build_workbook(store, tmp_path / "wb.xlsx", first_year="2024")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Dashboard", "Verträge & Abos", "Forecast",
                             "2024", "2025"]
    ws = wb["2024"]
    kpis = {ws.cell(row=4, column=2 + i * 2).value:
            ws.cell(row=5, column=2 + i * 2).value for i in range(5)}
    assert kpis["Einnahmen"] == 3000.0         # Retoure zählt NICHT als Einnahme
    assert kpis["Ausgaben"] == -70.0           # Shopping netto, ohne Neutrale
    assert kpis["Sondereffekte"] == 9999.0     # separat ausgewiesen
    # unkategorisierte Zeilen zählen mit (kein NULL-Trap); −40 tail
    # + 6 × −9,99 Netflix = −99,94
    ws25 = wb["2025"]
    kpis25 = {ws25.cell(row=4, column=2 + i * 2).value:
              ws25.cell(row=5, column=2 + i * 2).value for i in range(5)}
    assert kpis25["Ausgaben"] == -99.94


def test_abos_merge_namensvarianten_und_finden_jahresversicherungen():
    store = Store(":memory:")
    txs = []
    # Hallesche: Namensvariante wechselt (Umzug/FinTS-IBAN-Präfix) und der
    # Beitrag steigt — trotzdem EIN Vertrag, EINE Zeile, Betrag = aktuell
    for m in range(1, 13):
        txs.append(tx(booking_date=date(2024, m, 1), amount=Decimal("-80.00"),
                      counterparty="HALLESCHE Reinsburgstrasse 10",
                      category="Versicherung"))
    for m in range(1, 7):
        txs.append(tx(booking_date=date(2025, m, 1), amount=Decimal("-88.00"),
                      counterparty="DE24600501010001114445HALLESCHE",
                      category="Versicherung"))
    # Concordia: jährlich mit Beitragsanpassung → nie 3× derselbe Betrag,
    # muss über die Kategorie trotzdem auftauchen
    txs.append(tx(booking_date=date(2024, 3, 14), amount=Decimal("-277.20"),
                  counterparty="Concordia", category="Versicherung"))
    txs.append(tx(booking_date=date(2025, 2, 2), amount=Decimal("-264.00"),
                  counterparty="Concordia", category="Versicherung"))
    store.add(txs)

    abos = _detect_abos(store, today=date(2025, 7, 1))

    hallesche = [a for a in abos if "hallesche" in a["anbieter"].lower()]
    assert len(hallesche) == 1
    assert hallesche[0]["aktiv"]
    assert hallesche[0]["betrag"] == 8800
    assert hallesche[0]["seit"] == "2024-01-01"
    assert hallesche[0]["anbieter"] == "HALLESCHE"  # IBAN-Präfix entfernt

    concordia = [a for a in abos if a["anbieter"] == "Concordia"]
    assert len(concordia) == 1
    assert concordia[0]["turnus"] == "jährlich"
    assert concordia[0]["betrag"] == 26400
    assert concordia[0]["aktiv"]
